from docx import Document
from docx.shared import Pt


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)

    title = "基于Pcap的网络流量分析平台的设计与实现"
    doc.add_heading(title, level=0)

    def h(text, level=1):
        doc.add_heading(text, level=level)

    def p(text):
        doc.add_paragraph(text)

    h("摘要")
    p(
        "随着互联网应用规模持续增长，网络安全事件和数据泄露风险不断增加，网络流量分析已成为运维与安全检测中的关键技术手段。"
        "针对传统工具学习成本高、部署复杂等问题，本文设计并实现了一套基于Pcap文件的网络流量分析平台。系统采用Flask作为后端框架，"
        "Scapy作为流量解析引擎，SQLite作为数据存储，前端基于Bootstrap与ECharts实现可视化展示，实现了流量上传、数据包解析、协议统计、"
        "时间序列展示、Top IP分析、会话提取、敏感信息识别、攻击行为检测与PDF报告导出等核心功能。测试结果表明，系统能够稳定完成中小规模"
        "Pcap文件分析任务，并对SQL注入、目录遍历、弱口令暴露等行为给出有效告警。"
    )
    p("关键词：Pcap；网络流量分析；Flask；Scapy；可视化；安全检测")

    h("Abstract")
    p(
        "This thesis designs and implements a Pcap-based network traffic analysis platform. The system adopts Flask as the backend framework, "
        "Scapy for packet parsing, SQLite for data persistence, and Bootstrap with ECharts for visualization. It supports file upload, packet parsing, "
        "protocol statistics, timeline analysis, Top IP ranking, session extraction, sensitive information detection, attack behavior detection, "
        "and PDF report export. Experimental results show that the platform can stably process small and medium Pcap files and detect suspicious behaviors "
        "such as SQL injection and directory traversal."
    )
    p("Key Words: Pcap, Network Traffic Analysis, Flask, Scapy, Visualization, Security Detection")

    h("第1章 绪论")
    h("1.1 研究背景与意义", 2)
    p(
        "网络通信规模持续扩大，业务数字化程度不断提升，导致网络攻击面同步扩张。Pcap文件能够完整记录网络通信过程，"
        "是流量分析与安全取证的重要数据基础。设计一套轻量化、低门槛、可视化的流量分析平台，具有明显的工程价值与教学价值。"
    )
    h("1.2 国内外研究现状", 2)
    p(
        "国外在网络流量分析工具与规则引擎方面发展较早，形成了Wireshark、Snort、Suricata等成熟生态。"
        "国内相关研究在异常检测、可视化展示和场景化落地方面发展迅速，但在“低成本部署+易用性+可扩展性”三者平衡方面仍有改进空间。"
    )
    h("1.3 研究内容", 2)
    p("本文围绕“上传—解析—存储—分析—可视化—导出”闭环，完成需求分析、系统设计、功能实现与测试验证。")
    h("1.4 论文结构安排", 2)
    p("全文共7章：绪论、关键技术、需求分析、系统设计、系统实现、测试分析、总结展望。")

    h("第2章 关键技术与理论基础")
    h("2.1 Pcap与网络协议基础", 2)
    p("Pcap是抓包标准格式，数据包通常包含链路层、网络层、传输层及部分应用层数据。系统重点分析IP、TCP/UDP与HTTP相关字段。")
    h("2.2 Flask框架", 2)
    p("Flask轻量灵活，适合本科项目快速实现Web服务与接口。")
    h("2.3 Scapy库", 2)
    p("Scapy支持离线解析Pcap并提取包字段，是本系统核心解析组件。")
    h("2.4 SQLite与ORM", 2)
    p("SQLite部署简单，结合SQLAlchemy能够快速完成模型映射与数据持久化。")
    h("2.5 前端可视化技术", 2)
    p("Bootstrap构建页面结构，ECharts完成协议分布与时序曲线展示。")

    h("第3章 系统需求分析")
    h("3.1 功能需求", 2)
    p("系统支持pcap/cap上传、数据包解析、协议统计、时间序列分析、Top IP分析、会话提取、敏感信息识别、攻击检测及PDF导出。")
    h("3.2 非功能需求", 2)
    p("系统需满足易用性、性能、可靠性、可扩展性与安全性等要求。")

    h("第4章 系统设计")
    h("4.1 总体架构", 2)
    p("系统采用B/S三层架构：表示层、业务逻辑层、数据访问层。")
    h("4.2 模块设计", 2)
    p("包含文件上传模块、解析模块、统计分析模块、会话模块、安全检测模块、可视化模块与报告导出模块。")
    h("4.3 数据库设计", 2)
    p("主要数据表包括CaptureFile、PacketInfo、SessionInfo、SensitiveInfo、AttackInfo。")

    h("第5章 系统实现")
    h("5.1 开发环境", 2)
    p("Windows 10、Python 3.10、Flask、Scapy、SQLite、Bootstrap、ECharts。")
    h("5.2 核心实现说明", 2)
    p("上传后系统生成file_uuid并保存文件；解析模块提取包字段入库；会话模块按连接聚合内容；检测模块通过正则识别敏感信息与攻击特征；结果页展示统计图表并支持PDF导出。")

    h("第6章 系统测试与结果分析")
    h("6.1 测试方案", 2)
    p("从功能测试、异常测试和样本测试三个维度验证系统。")
    h("6.2 测试结果", 2)
    p("系统可稳定完成中小规模Pcap分析，敏感信息与攻击检测对测试样本命中有效，结果可视化清晰，导出报告完整。")
    h("6.3 存在问题", 2)
    p("当前以离线分析为主，协议深度解析与实时处理能力仍有提升空间。")

    h("第7章 总结与展望")
    h("7.1 总结", 2)
    p("本文完成了基于Pcap的网络流量分析平台设计与实现，形成了可运行、可展示、可扩展的毕设成果。")
    h("7.2 展望", 2)
    p("后续可扩展实时抓包、更多协议解析、IP地理可视化、规则引擎与权限审计等功能。")

    h("参考文献")
    p("[1] Combs G. Wireshark User's Guide.")
    p("[2] Biondi P. Scapy Documentation.")
    p("[3] Grinberg M. Flask Web Development.")
    p("[4] Stallings W. Data and Computer Communications.")

    out = r"F:\work\cursor\基于Pcap的网络流量分析平台的设计与实现-论文初稿.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
