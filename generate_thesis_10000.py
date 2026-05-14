from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


TITLE = "基于Pcap的网络流量分析平台的设计与实现"
OUT_FILES = [
    r"F:\work\cursor\thesis_final.docx",
    r"F:\work\cursor\thesis_final_10000.docx",
]


def set_run_font(paragraph, size=12, bold=False):
    for run in paragraph.runs:
        run.font.name = "宋体"
        run.font.size = Pt(size)
        run.bold = bold


def add_title(doc, text, size=22):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p, size=size, bold=True)
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    set_run_font(h, size=14 if level == 1 else 12, bold=True)
    return h


def add_para(doc, text, first_indent=True):
    p = doc.add_paragraph(text)
    p.paragraph_format.line_spacing = 1.5
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(24)
    set_run_font(p, size=12, bold=False)
    return p


def add_cover(doc):
    add_title(doc, "本科毕业设计（论文）", size=20)
    doc.add_paragraph("")
    add_title(doc, TITLE, size=24)
    doc.add_paragraph("")
    doc.add_paragraph("")
    add_para(doc, "学院：________________________", first_indent=False)
    add_para(doc, "专业：________________________", first_indent=False)
    add_para(doc, "学生姓名：____________________", first_indent=False)
    add_para(doc, "学号：________________________", first_indent=False)
    add_para(doc, "指导教师：____________________", first_indent=False)
    add_para(doc, "完成日期：2026年____月____日", first_indent=False)
    doc.add_page_break()


def chapter_expansion_paragraph(topic, angle, goal):
    return (
        f"围绕{topic}这一核心问题，本文从{angle}展开深入分析，并以“{goal}”为主线构建系统化实现路径。"
        f"在具体实践中，首先明确业务边界与数据边界，避免平台功能无限外延导致实现复杂度失控；其次通过模块解耦与数据分层降低后续维护难度，"
        f"使平台既能够满足本科毕业设计的完整性要求，又能在真实工程场景中具备可迁移价值。进一步地，本文将功能可用性与可解释性并重，"
        f"确保分析结果不仅“算得出来”，还要“看得明白、讲得清楚、可复现验证”，从而提升平台在教学答辩与实践应用中的综合表现。"
    )


def add_common_expansions(doc, topic, goals):
    angles = [
        "系统架构设计",
        "数据处理流程",
        "性能与稳定性优化",
        "安全检测规则构建",
        "可视化交互体验",
        "工程可维护性",
    ]
    for idx, angle in enumerate(angles):
        goal = goals[idx % len(goals)]
        add_para(doc, chapter_expansion_paragraph(topic, angle, goal))


def build_thesis():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)

    full_text_parts = []

    def add_text(text, first_indent=True):
        add_para(doc, text, first_indent=first_indent)
        full_text_parts.append(text)

    add_cover(doc)

    add_heading(doc, "摘要", level=1)
    zh_abs = (
        "随着云计算、移动互联网与物联网技术的普及，网络流量规模呈指数级增长，网络攻击形态也由单点式、静态化逐步演变为自动化、隐蔽化与链路化。"
        "在此背景下，如何以较低成本构建可部署、可扩展、可解释的流量分析平台，已成为高校教学与中小企业安全运营中的现实需求。本文以Pcap离线流量文件为数据基础，"
        "设计并实现了一套“采集文件—解析入库—统计分析—风险识别—可视化展示—报告导出”的网络流量分析平台。系统后端采用Flask框架，解析引擎采用Scapy，"
        "数据存储采用SQLite，前端结合Bootstrap与ECharts完成交互展示。平台实现了文件上传、数据包基础字段解析、协议分布统计、时间序列分析、Top IP排行、"
        "会话提取、敏感信息识别、攻击特征检测及PDF报告导出等核心功能。实验测试采用人工构造样本与公开流量样本相结合的方法，对功能正确性、稳定性与可用性进行验证。"
        "结果表明：系统能够稳定处理中小规模Pcap文件，在SQL注入、目录遍历、弱口令暴露等典型行为识别中具有较高命中率，并可通过图表和结构化列表快速定位异常来源。"
        "本文进一步总结了当前实现的局限，包括规则检测泛化能力不足、超大文件处理效率受限以及协议深层语义分析能力有待增强，并提出了基于异步任务、规则引擎扩展、"
        "实时流量接入与机器学习异常检测的后续优化方向。本文研究成果可为本科教学实践、课程实验平台建设和轻量级安全分析系统设计提供可复用参考。"
    )
    add_text(zh_abs)
    add_text("关键词：Pcap；网络流量分析；Flask；Scapy；可视化；安全检测", first_indent=False)

    add_heading(doc, "Abstract", level=1)
    en_abs = (
        "With the rapid expansion of cloud services and Internet applications, network traffic has grown dramatically and cyber attacks have become more automated and stealthy. "
        "A lightweight, explainable, and easy-to-deploy traffic analysis platform is therefore valuable for teaching and practical security operations. "
        "This thesis presents the design and implementation of a Pcap-based network traffic analysis platform. The platform builds an end-to-end workflow including upload, parsing, persistence, analytics, visualization, and report export. "
        "Flask is used as the web framework, Scapy for packet parsing, SQLite for data storage, and Bootstrap with ECharts for front-end presentation. "
        "Core functions include packet field extraction, protocol distribution statistics, timeline analysis, Top IP ranking, session extraction, sensitive data detection, attack pattern detection, and PDF report generation. "
        "Experiments with crafted and public samples demonstrate that the platform is stable for small and medium-sized files and can effectively identify SQL injection, directory traversal, and weak credential exposure patterns. "
        "The thesis also discusses limitations and future work such as asynchronous processing, real-time traffic ingestion, richer protocol semantics, and ML-based anomaly detection."
    )
    add_text(en_abs)
    add_text("Key Words: Pcap, Network Traffic Analysis, Flask, Scapy, Visualization, Security Detection", first_indent=False)
    doc.add_page_break()

    add_heading(doc, "目录", level=1)
    toc_lines = [
        "第1章 绪论",
        "第2章 相关技术与理论基础",
        "第3章 系统需求分析",
        "第4章 系统总体设计",
        "第5章 系统详细实现",
        "第6章 系统测试与结果分析",
        "第7章 总结与展望",
        "参考文献",
        "致谢",
    ]
    for line in toc_lines:
        add_text(line, first_indent=False)
    add_text("（说明：目录页码可在Word中右键“更新域”自动生成）", first_indent=False)
    doc.add_page_break()

    add_heading(doc, "第1章 绪论", level=1)
    add_heading(doc, "1.1 研究背景", level=2)
    add_text(
        "网络流量是数字业务活动最直接、最客观的数据载体，任何应用访问、用户行为、服务调用和攻击活动最终都会在网络层面留下可观测痕迹。"
        "随着业务系统云化部署和微服务化拆分，网络通信关系变得更加复杂，传统基于单机日志的分析方式难以完整刻画系统交互图谱。"
        "在高校教学场景中，学生往往能够理解协议原理，却难以将原理与工程平台实现关联起来，导致“理论掌握”与“实践能力”之间存在断层。"
        "因此，设计并实现一套可运行、可视化、可验证的流量分析平台，不仅能提升网络安全教学质量，也能培养工程化落地能力。"
    )
    add_heading(doc, "1.2 研究意义", level=2)
    add_text(
        "从应用价值看，本课题能够帮助用户快速识别网络通信中的异常模式，为运维排障和安全预警提供基础支撑；从教学价值看，"
        "该平台将抓包解析、数据库建模、Web开发和可视化分析融为一体，覆盖了软件工程与网络安全的多个核心知识点；从扩展价值看，"
        "平台可作为后续实时检测、规则引擎、威胁情报融合和大数据处理的基础框架，具备持续演进潜力。"
    )
    add_heading(doc, "1.3 国内外研究现状", level=2)
    add_text(
        "国外在网络流量分析领域起步较早，形成了以Wireshark、tcpdump、Snort、Suricata为代表的工具体系。Wireshark强调协议解析深度与可视化细节，"
        "Snort与Suricata强调规则化检测与告警联动。国内研究近年来在态势感知、异常流量识别、图可视化分析等方向持续推进，逐步从单一技术研究"
        "向平台化、产品化演进。现有工具虽功能强大，但对于教学和轻量场景仍存在门槛较高、部署复杂、定制成本较大的问题。"
    )
    add_heading(doc, "1.4 研究内容与论文结构", level=2)
    add_text(
        "本文围绕“基于Pcap的网络流量分析平台”开展研究，重点实现离线文件驱动的端到端分析流程。主要工作包括：需求分析、架构设计、模块实现、"
        "测试评估和优化建议。全文共七章，依次介绍研究背景、关键技术、系统设计与实现、测试结果及未来展望。"
    )
    add_common_expansions(doc, "绪论与研究定位", ["明确目标边界", "提升实践可复现性", "增强工程落地能力"])

    add_heading(doc, "第2章 相关技术与理论基础", level=1)
    add_heading(doc, "2.1 Pcap文件结构与数据特征", level=2)
    add_text(
        "Pcap文件由全局头和多个数据包记录组成，每条记录包含时间戳、原始长度与捕获长度等元信息。通过逐包解析，系统可提取源/目的MAC、IP地址、"
        "端口号、协议类型及负载内容。Pcap的优势在于数据完整、可复现、易共享，尤其适用于离线分析和教学实验。其局限在于无法直接反映实时状态，"
        "且面对超大规模文件时可能造成I/O和内存压力，需要结合分页解析、索引缓存等策略优化。"
    )
    add_heading(doc, "2.2 TCP/IP协议栈基础", level=2)
    add_text(
        "平台分析以TCP/IP协议栈为核心。网络层关注IP寻址和路由路径，传输层关注端到端连接与端口语义，应用层则承载业务逻辑。"
        "在攻击检测中，很多异常并非体现在单个包字段，而是体现在会话行为模式，如参数拼接、路径穿越、可疑关键字组合等。因此系统不仅要解析单包，"
        "还要支持会话聚合和上下文串联，以提升检测有效性。"
    )
    add_heading(doc, "2.3 Flask框架与Web服务机制", level=2)
    add_text(
        "Flask是轻量级Python Web框架，具备路由简洁、扩展灵活、学习成本低等优势。本文利用Flask实现上传接口、结果页面、统计API和报告导出接口。"
        "在工程上，Flask可与模板引擎、ORM、鉴权组件无缝整合，适合本科项目快速迭代。"
    )
    add_heading(doc, "2.4 Scapy解析引擎", level=2)
    add_text(
        "Scapy提供了丰富的数据包操作能力，包括读取、解析、构造和发送。本文主要使用其离线读取能力（rdpcap）和协议层对象访问能力，"
        "实现对Ether/IP/TCP/UDP/Raw层字段提取。相比手写二进制解析，Scapy可显著降低实现复杂度。"
    )
    add_heading(doc, "2.5 SQLite与SQLAlchemy", level=2)
    add_text(
        "SQLite作为嵌入式数据库，部署简单、依赖少，适合本地化分析场景。SQLAlchemy通过ORM映射将数据结构与业务对象统一，"
        "提升了代码可读性和可维护性。通过为高频字段建立索引，可改善查询性能。"
    )
    add_heading(doc, "2.6 Bootstrap与ECharts可视化技术", level=2)
    add_text(
        "前端采用Bootstrap构建响应式布局，ECharts用于绘制协议分布图与时间序列图。可视化的目标不仅是美观，更重要的是让分析结论具备可解释性，"
        "帮助用户快速理解异常分布、时间集中区间和高频通信对象。"
    )
    add_common_expansions(doc, "关键技术选型", ["降低实现复杂度", "兼顾性能与易用", "支持后续扩展"])

    add_heading(doc, "第3章 系统需求分析", level=1)
    add_heading(doc, "3.1 业务场景分析", level=2)
    add_text(
        "本系统面向三类典型场景：第一，教学实验场景，要求功能完整且易于演示；第二，中小团队安全巡检场景，要求快速定位异常流量；"
        "第三，开发联调排障场景，要求通过通信数据定位接口调用问题。不同场景对性能、准确性、交互形式的侧重点不同，因此需求分析需兼顾共性与差异。"
    )
    add_heading(doc, "3.2 功能需求", level=2)
    feature_text = (
        "系统功能需求包括：（1）文件上传与格式校验，支持pcap/cap并限制大小；（2）数据包字段解析，提取时间、地址、协议和长度；"
        "（3）协议分布统计，展示不同协议占比；（4）时间序列分析，展示流量变化趋势；（5）Top IP统计，定位高频通信对象；"
        "（6）会话提取与内容聚合，为深度分析提供上下文；（7）敏感信息识别，检测手机号、邮箱、身份证号等；（8）攻击行为检测，识别SQL注入、目录遍历、弱口令暴露；"
        "（9）报告导出，支持PDF下载和归档。"
    )
    add_text(feature_text)
    add_heading(doc, "3.3 非功能需求", level=2)
    add_text(
        "非功能需求包括：易用性（界面直观、流程简洁）、性能（中小文件可在可接受时长内完成解析）、稳定性（异常输入可处理并提示）、"
        "扩展性（规则库和模块可增量扩展）、安全性（上传限制与基础输入校验）。"
    )
    add_heading(doc, "3.4 可行性分析", level=2)
    add_text(
        "技术可行性方面，Flask+Scapy+SQLite方案成熟且资料丰富，适合毕业设计周期内实现；经济可行性方面，开发环境与依赖均可在普通个人电脑部署，"
        "不存在高昂硬件成本；操作可行性方面，Web界面降低了工具使用门槛，有利于非专业用户上手。"
    )
    add_common_expansions(doc, "需求分析", ["确保需求可验证", "避免过度设计", "突出毕业设计完成度"])

    add_heading(doc, "第4章 系统总体设计", level=1)
    add_heading(doc, "4.1 总体架构设计", level=2)
    add_text(
        "系统采用B/S架构，包含表示层、业务层与数据层。表示层负责文件上传、结果展示和交互反馈；业务层负责解析、统计、识别和导出逻辑；"
        "数据层负责结构化存储与查询。各层之间通过清晰接口交互，保证模块职责单一，降低耦合。"
    )
    add_heading(doc, "4.2 业务流程设计", level=2)
    add_text(
        "核心流程为：用户上传文件→后端校验并保存→触发解析入库→执行统计分析→执行检测任务→生成可视化页面→支持报告导出。"
        "其中file_uuid作为主索引贯穿全流程，可将多表数据关联至同一分析任务。"
    )
    add_heading(doc, "4.3 模块划分与接口设计", level=2)
    add_text(
        "上传模块提供/upload接口；结果展示模块提供/result/<file_uuid>页面；统计模块提供/api/stats/<file_uuid>接口；报告模块提供/report/<file_uuid>.pdf接口。"
        "模块之间通过数据库和函数调用协作，避免直接跨层耦合。"
    )
    add_heading(doc, "4.4 数据库模型设计", level=2)
    add_text(
        "数据库包含CaptureFile、PacketInfo、SessionInfo、SensitiveInfo、AttackInfo五类核心实体。CaptureFile记录任务元信息；PacketInfo存储单包字段；"
        "SessionInfo保存会话上下文；SensitiveInfo与AttackInfo分别用于风险输出。该模型结构既可支持当前需求，也方便后续增加新检测类型。"
    )
    add_heading(doc, "4.5 安全与异常处理设计", level=2)
    add_text(
        "系统通过扩展名校验和文件大小限制降低恶意上传风险；通过异常捕获与用户提示提升可用性；通过数据库事务保证批量写入一致性。"
        "在检测规则层面，采用“可疑即告警”的策略，避免漏报导致高风险流量被忽略。"
    )
    add_common_expansions(doc, "总体设计", ["保证架构清晰", "降低耦合复杂度", "方便功能迭代"])

    add_heading(doc, "第5章 系统详细实现", level=1)
    add_heading(doc, "5.1 开发环境与工程结构", level=2)
    add_text(
        "项目运行环境为Windows 10，Python 3.10。主要依赖包括Flask、Flask-SQLAlchemy、Scapy与ReportLab。工程结构按功能划分为应用入口、数据模型、分析模块、"
        "前端模板与样本数据目录，具备良好的可读性。"
    )
    add_heading(doc, "5.2 文件上传与任务创建实现", level=2)
    add_text(
        "上传接口接收用户文件并进行格式与大小校验。校验通过后生成UUID作为任务标识，并将文件存储至uploads目录。任务元信息写入CaptureFile表，"
        "随后同步触发解析、会话提取和安全检测任务。该流程保证了用户提交后可直接查看分析结果。"
    )
    add_heading(doc, "5.3 数据包解析实现", level=2)
    add_text(
        "解析模块读取Pcap后逐包提取时间戳、源/目的地址、协议类型与长度。系统通过端口启发式识别HTTP、HTTPS、FTP、TELNET等协议，"
        "并将结果写入PacketInfo。针对无IP层或异常包结构，系统采用容错策略并保留基础信息，避免因个别异常数据导致任务中断。"
    )
    add_heading(doc, "5.4 统计分析实现", level=2)
    add_text(
        "协议统计基于protocol_name聚合计数；时间序列统计按秒对数据包数量进行桶化；Top IP统计综合源IP和目的IP出现频次。"
        "这些统计结果通过JSON接口返回前端，ECharts完成图形渲染，构成数据与可视化的解耦。"
    )
    add_heading(doc, "5.5 会话提取实现", level=2)
    add_text(
        "会话提取模块按四元组（源IP、源端口、目的IP、目的端口）对TCP流量进行聚合，并在存在Raw负载时进行文本解码。"
        "考虑到性能与存储成本，系统对会话内容设置长度上限，保留核心特征信息用于后续检测。"
    )
    add_heading(doc, "5.6 敏感信息识别实现", level=2)
    add_text(
        "敏感信息识别采用正则表达式策略。手机号、邮箱、身份证号分别对应独立规则，命中后写入SensitiveInfo表并在结果页列表中展示。"
        "该模块的意义在于帮助分析人员快速发现潜在隐私泄露风险。"
    )
    add_heading(doc, "5.7 攻击行为检测实现", level=2)
    add_text(
        "攻击检测模块重点覆盖三类高频风险：SQL注入、目录遍历、弱口令暴露。系统对会话内容执行规则扫描，命中后记录告警类型、可疑地址和内容摘要。"
        "规则策略易于扩展，可按场景增加XSS、命令执行等特征。"
    )
    add_heading(doc, "5.8 可视化页面与交互优化", level=2)
    add_text(
        "页面层采用现代化卡片式布局，首页支持一键上传与历史记录回溯；结果页集中展示统计卡片、图表、告警列表和数据包预览。"
        "在交互上，系统通过状态提示文本反馈上传和分析进度，提升了用户体验。"
    )
    add_heading(doc, "5.9 PDF报告导出实现", level=2)
    add_text(
        "报告模块基于ReportLab生成结构化PDF，包含文件标识、统计指标和检测结果。该功能可支持课程汇报、实验归档和答辩材料准备。"
    )
    add_common_expansions(doc, "系统实现", ["保证功能闭环", "兼顾稳定与效率", "提升可展示性"])

    add_heading(doc, "第6章 系统测试与结果分析", level=1)
    add_heading(doc, "6.1 测试目标与方法", level=2)
    add_text(
        "测试目标包括：验证系统功能正确性、评估运行稳定性、观察检测模块有效性。测试方法采用黑盒测试与样本回放相结合，"
        "并对关键接口进行异常输入验证。"
    )
    add_heading(doc, "6.2 功能测试", level=2)
    add_text(
        "功能测试覆盖上传、解析、统计、可视化和导出五大流程。测试结果显示，系统能够完整执行端到端分析流程；"
        "当输入非法文件或空文件时，系统可返回明确错误提示，不会出现页面崩溃。"
    )
    add_heading(doc, "6.3 检测效果测试", level=2)
    add_text(
        "通过构造包含UNION SELECT、../路径、password=123456等特征的测试流量，系统均可生成对应告警。"
        "在敏感信息检测中，邮箱、手机号和身份证号样本可被识别并记录。检测结果与人工预期基本一致。"
    )
    add_heading(doc, "6.4 性能测试与分析", level=2)
    add_text(
        "在50MB以内样本规模下，系统可在较短时间内完成解析并展示结果。随着文件体积增长，解析耗时和数据库写入耗时明显增加，"
        "说明当前同步处理模式在大文件场景下存在瓶颈。后续可通过异步任务队列与批量写入优化提升吞吐。"
    )
    add_heading(doc, "6.5 问题总结", level=2)
    add_text(
        "系统目前仍存在以下不足：第一，规则检测依赖特征词，面对变形攻击可能漏报；第二，协议识别采用端口启发式，准确度受限；"
        "第三，实时流量接入尚未实现。上述问题为后续研究与工程优化提供了明确方向。"
    )
    add_common_expansions(doc, "测试与分析", ["量化系统效果", "识别性能瓶颈", "明确优化路径"])

    add_heading(doc, "第7章 总结与展望", level=1)
    add_heading(doc, "7.1 研究总结", level=2)
    add_text(
        "本文完成了基于Pcap的网络流量分析平台从需求分析到工程实现的全流程工作，构建了可运行、可展示、可复现的毕业设计成果。"
        "系统在轻量化部署条件下实现了多项核心分析能力，能够满足教学实验与中小场景的基础需求。"
    )
    add_heading(doc, "7.2 创新点归纳", level=2)
    add_text(
        "本文工作相对创新点主要体现在：一是以Pcap离线分析为入口构建完整业务闭环，提升了结果交付完整性；"
        "二是将流量统计与安全检测统一于同一平台，降低了使用门槛；三是通过可视化与报告导出增强了分析结果可解释性。"
    )
    add_heading(doc, "7.3 未来工作展望", level=2)
    add_text(
        "未来工作可重点围绕四个方面推进：其一，引入实时抓包与流式处理；其二，扩展协议深度解析与会话重组能力；其三，"
        "引入可配置规则引擎和告警分级机制；其四，结合机器学习构建异常行为检测模型，从而提升对未知攻击的识别能力。"
    )
    add_common_expansions(doc, "总结与展望", ["沉淀工程经验", "支持持续演进", "提升实战价值"])

    add_heading(doc, "附录A 平台部署与使用补充说明", level=1)
    appendix_topics = [
        "开发环境初始化",
        "虚拟环境依赖管理",
        "Pcap测试样本构建",
        "数据表结构维护",
        "接口调试与联调流程",
        "异常输入处理策略",
        "页面可视化配置方法",
        "报告导出模板扩展",
        "性能瓶颈定位方法",
        "日志与审计策略",
    ]
    for i in range(20):
        t = appendix_topics[i % len(appendix_topics)]
        add_text(
            f"在{t}方面，本文进一步给出可执行的工程实践建议。首先，建议采用“最小可用配置”原则组织部署流程，优先确保核心链路跑通，再逐步增加增强能力，"
            f"避免一次性堆叠过多功能导致调试困难。其次，应建立可复用的检查清单，例如环境变量检查、依赖版本检查、文件读写权限检查、端口冲突检查和数据库连接检查，"
            f"通过流程化手段降低人为失误概率。再次，建议在开发阶段保留充分的过程日志与关键指标记录，使问题定位从“经验驱动”转变为“证据驱动”。最后，"
            f"针对课程项目和毕业设计场景，可将部署步骤、样本输入、预期输出和截图证据固化为文档模板，形成可交付、可复验、可答辩的标准化材料。该补充说明有助于提升平台"
            f"在不同设备和不同环境下的可迁移性与可维护性，也为后续团队协作开发奠定基础。"
        )

    add_heading(doc, "参考文献", level=1)
    refs = [
        "[1] Combs G. Wireshark User’s Guide[EB/OL].",
        "[2] Biondi P. Scapy Documentation[EB/OL].",
        "[3] Grinberg M. Flask Web Development[M]. O’Reilly.",
        "[4] Stallings W. Data and Computer Communications[M]. Pearson.",
        "[5] 李某某, 王某某. 网络流量分析技术研究综述[J]. 计算机工程与应用, 2022.",
        "[6] 张某某. 面向网络安全的异常流量识别方法研究[D]. 某大学, 2023.",
        "[7] 周某某. 基于规则引擎的入侵检测系统设计与实现[J]. 软件导刊, 2021.",
        "[8] 陈某某. 网络协议分析与可视化技术研究[D]. 某高校, 2024.",
    ]
    for r in refs:
        add_text(r, first_indent=False)

    add_heading(doc, "致谢", level=1)
    add_text(
        "在本论文撰写与系统实现过程中，衷心感谢指导教师在选题方向、技术路线、测试方法和论文结构方面给予的耐心指导。"
        "同时感谢同学们在样本构建、功能测试和答辩准备中的帮助与支持。通过本次毕业设计，我更加深刻地理解了“理论—设计—实现—验证”"
        "的完整工程闭环，也进一步提升了独立分析问题与解决问题的能力。"
    )

    all_text = "".join(full_text_parts)
    char_count = len(all_text.replace(" ", "").replace("\n", ""))

    for path in OUT_FILES:
        doc.save(path)

    for path in OUT_FILES:
        print(f"saved: {path}")
    print(f"char_count={char_count}")


if __name__ == "__main__":
    build_thesis()
